from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml

doc = Document()

# Judul utama
title = doc.add_heading('LAPORAN LATIHAN INDIVIDU - PEMROGRAMAN WEB PHP', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Info mahasiswa
info = doc.add_paragraph()
info.alignment = WD_ALIGN_PARAGRAPH.CENTER
info.add_run('Mata Kuliah: Pemrograman Web\n').bold = True
info.add_run('Tugas: Latihan Individu - Kartu Hasil Studi (KHS)\n')
info.add_run('Deadline: 01 November 2026 Pukul 23:59 WIB')

doc.add_paragraph()

# Link GitHub
p_link = doc.add_paragraph()
p_link.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_link = p_link.add_run('Link Repositori GitHub: https://github.com/a3snewbaru990-ui/tugas-pemweb-php')
run_link.italic = True
run_link.font.color.rgb = RGBColor(0x1a, 0x23, 0x7e)

doc.add_paragraph()

# Screenshot output
p_ss = doc.add_paragraph()
p_ss.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_ss = p_ss.add_run('Screenshot Output di Localhost (http://localhost/index.php):')
run_ss.bold = True
run_ss.font.size = Pt(12)

p_note = doc.add_paragraph()
p_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_note = p_note.add_run('[Sisipkan screenshot output localhost di sini]')
run_note.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
run_note.italic = True

doc.add_paragraph()

# Tabel nilai
doc.add_heading('Hasil Nilai Mata Kuliah', level=1)

data = [
    ['Nama Mata Kuliah', 'SKS', 'Nilai Angka', 'Nilai Huruf', 'Status'],
    ['Pemrograman Web', '3', '88', 'A', 'LULUS'],
    ['Basis Data', '3', '72', 'B', 'LULUS'],
    ['Jaringan Komputer', '2', '65', 'C', 'LULUS'],
    ['Sistem Operasi', '3', '55', 'C', 'LULUS'],
    ['Pemrograman Mobile', '4', '40', 'D', 'MENGULANG'],
    ['Kecerdasan Buatan', '3', '30', 'E', 'MENGULANG'],
]

table = doc.add_table(rows=len(data), cols=len(data[0]))
table.style = 'Table Grid'
table.alignment = WD_TABLE_ALIGNMENT.CENTER

header_colors = {
    'LULUS': RGBColor(0x2e, 0x7d, 0x32),
    'MENGULANG': RGBColor(0xc6, 0x28, 0x28),
}

for i, row in enumerate(data):
    for j, val in enumerate(row):
        cell = table.cell(i, j)
        cell.text = val
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(11)
                if i == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(0xff, 0xff, 0xff)
                    shd = parse_xml(r'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="clear" w:fill="1A237E"/>')
                    cell._tc.get_or_add_tcPr().append(shd)
                elif j == 4:
                    run.bold = True
                    run.font.color.rgb = header_colors.get(val, RGBColor(0, 0, 0))
                    if val == 'MENGULANG':
                        shd = parse_xml(r'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="clear" w:fill="FFEBEE"/>')
                        cell._tc.get_or_add_tcPr().append(shd)

doc.add_paragraph()

# Ringkasan IPS
doc.add_heading('Ringkasan IPS', level=1)

ringkasan = [
    ('Total SKS diambil:', '18 SKS'),
    ('IPS Anda:', '1.94'),
    ('Batas Maksimal SKS Semester:', '20 SKS (IPS < 3.00 maksimal 20 SKS)'),
]

for label, value in ringkasan:
    p = doc.add_paragraph()
    p.add_run(label + ' ').bold = True
    p.add_run(value)

doc.add_paragraph()

# Penjelasan program
doc.add_heading('Penjelasan Program', level=1)
poin = [
    'Array multidimensi $matkul berisi 6 mata kuliah dengan kunci: nama, sks, dan nilai (skala 0-100).',
    'Fungsi hitungBobot($nilai) mengembalikan array berisi huruf (A, B, C, D, E) dan bobot (4.0, 3.0, 2.0, 1.0, 0.0).',
    'Status LULUS jika bobot >= 2.0, MENGULANG jika bobot < 2.0.',
    'Baris mata kuliah yang statusnya MENGULANG ditandai dengan warna merah.',
    'IPS dihitung dari total bobot x SKS dibagi total SKS.',
    'Batas maksimal SKS: IPS >= 3.00 boleh ambil 24 SKS, IPS < 3.00 maksimal 20 SKS.',
]
for p_text in poin:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(p_text)

doc.add_paragraph()

# Kesimpulan
doc.add_heading('Kesimpulan', level=1)
p_kesimpulan = doc.add_paragraph()
p_kesimpulan.add_run(
    'Dari hasil perhitungan, IPS yang diperoleh adalah 1.94 dengan total SKS 18. '
    'Karena IPS kurang dari 3.00, batas maksimal SKS semester berikutnya adalah 20 SKS. '
    'Terdapat 2 mata kuliah yang statusnya MENGULANG (Pemrograman Mobile dan Kecerdasan Buatan).'
)

output_path = r'D:\python\Laporan_Latihan_Individu_Pemweb.docx'
doc.save(output_path)
print(f'File berhasil disimpan: {output_path}')